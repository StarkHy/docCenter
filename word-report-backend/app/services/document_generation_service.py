import os
import json
import uuid
from typing import Optional, List, Dict, Any
from datetime import datetime
from docxtpl import DocxTemplate
from docx.shared import Inches
from docx import Document as DocxDocument
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.models import DocumentHistory
from app.services.storage_service import storage_service
from app.repositories.template_repository import TemplateRepository
from app.repositories.datasource_repository import DataSourceRepository
from app.repositories.mapping_repository import MappingRepository
from app.repositories.document_history_repository import DocumentHistoryRepository
from app.services.chart_service import chart_service


class DocumentGenerationService:
    def __init__(
        self,
        template_repo: TemplateRepository,
        datasource_repo: DataSourceRepository,
        mapping_repo: MappingRepository,
        history_repo: DocumentHistoryRepository
    ):
        self.output_path = settings.WORD_OUTPUT_PATH
        self.template_repo = template_repo
        self.datasource_repo = datasource_repo
        self.mapping_repo = mapping_repo
        self.history_repo = history_repo
        os.makedirs(self.output_path, exist_ok=True)

    async def _fetch_datasource_data(self, datasource_id: int, session: AsyncSession, runtime_params: Optional[dict] = None) -> dict:
        from app.services.datasource_service import data_source_service
        runtime_params = runtime_params or {}
        try:
            datasource = await data_source_service.get_datasource(datasource_id, session)
            if not datasource:
                raise ValueError(f"数据源 {datasource_id} 不存在")

            ds_type = datasource.get("type")
            config = datasource.get("config", {})

            if ds_type == "http_api":
                return await self._fetch_http_api(config, runtime_params)
            elif ds_type == "json":
                return self._fetch_json(config)
            elif ds_type == "form":
                return config.get("fields", {}) if config else {}
            elif ds_type == "kingbase":
                return await self._fetch_kingbase(config, runtime_params)
            elif ds_type == "postgresql":
                return await self._fetch_postgresql(config, runtime_params)
            elif ds_type == "mysql":
                return await self._fetch_mysql(config, runtime_params)
            elif ds_type == "oracle":
                return await self._fetch_oracle(config, runtime_params)
            elif ds_type == "sql":
                return await self._fetch_sql(config, runtime_params, session)
            else:
                raise ValueError(f"不支持的数据源类型：{ds_type}")
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"获取数据源 {datasource_id} 失败：{str(e)}")

    async def _fetch_multiple_datasources(self, datasource_ids: List[int], session: AsyncSession, runtime_params: Optional[dict] = None) -> Dict[int, dict]:
        result = {}
        for ds_id in datasource_ids:
            try:
                data = await self._fetch_datasource_data(ds_id, session, runtime_params)
                result[ds_id] = data
            except ValueError:
                raise
            except Exception as e:
                raise ValueError(f"获取数据源 {ds_id} 失败：{str(e)}")
        return result

    async def _fetch_http_api(self, config: dict, runtime_params: Optional[dict] = None) -> dict:
        import requests
        runtime_params = runtime_params or {}
        try:
            method = config.get("method", "GET").upper()
            url = config.get("url", "")
            headers = config.get("headers", {})
            params = config.get("params", {})
            body_str = config.get("body", "")
            timeout = config.get("timeout", 30)
            parameters = config.get("parameters", [])
            
            for p in parameters:
                p_name = p.get("name", "")
                if p_name and p_name in runtime_params:
                    p["default"] = runtime_params[p_name]

            if not url:
                raise ValueError("请配置接口地址")

            url = self._substitute_sql_params(url, parameters)
            headers = self._substitute_params_in_dict(headers, parameters)
            params = self._substitute_params_in_dict(params, parameters)

            body = None
            if body_str and isinstance(body_str, str):
                try:
                    body = json.loads(body_str)
                    body = self._substitute_params_in_dict(body, parameters)
                except:
                    body = self._substitute_sql_params(body_str, parameters)
            elif isinstance(body_str, dict):
                body = self._substitute_params_in_dict(body_str, parameters)

            if method == "GET":
                response = requests.get(url, headers=headers, params=params, timeout=timeout)
            elif method == "POST":
                if isinstance(body, dict):
                    response = requests.post(url, headers=headers, params=params, json=body, timeout=timeout)
                else:
                    response = requests.post(url, headers=headers, params=params, data=body, timeout=timeout)
            else:
                raise ValueError(f"不支持的请求方法：{method}")

            response.raise_for_status()

            content_type = response.headers.get("content-type", "")
            if "application/json" in content_type:
                return response.json()
            else:
                return {"text": response.text}
        except requests.exceptions.Timeout:
            raise ValueError("API 请求超时，请检查网络或增加超时时间")
        except requests.exceptions.ConnectionError:
            raise ValueError("API 连接失败，请检查网络或 URL 是否正确")
        except requests.exceptions.HTTPError as e:
            raise ValueError(f"API 请求失败：HTTP {e.response.status_code}")
        except requests.exceptions.RequestException as e:
            raise ValueError(f"API 请求失败：{str(e)}")

    def _fetch_json(self, config: dict) -> dict:
        try:
            json_text = config.get("json_text", "")
            if not json_text:
                return {}
            return json.loads(json_text)
        except json.JSONDecodeError as e:
            raise ValueError(f"JSON 解析失败，请检查 JSON 格式：{str(e)}")
        except Exception as e:
            raise ValueError(f"JSON 数据处理失败：{str(e)}")

    async def _fetch_kingbase(self, config: dict, runtime_params: Optional[dict] = None) -> dict:
        import psycopg2
        runtime_params = runtime_params or {}
        try:
            host = config.get("host", "")
            port = config.get("port", "54321")
            database = config.get("database", "")
            user = config.get("user", "")
            password = config.get("password", "")
            sql = config.get("sql", "") or config.get("test_sql", "SELECT 1")
            params = config.get("params", [])
            
            for p in params:
                p_name = p.get("name", "")
                if p_name and p_name in runtime_params:
                    p["default"] = runtime_params[p_name]

            sql = self._substitute_sql_params(sql, params)
            
            conn = psycopg2.connect(
                host=host,
                port=port,
                database=database,
                user=user,
                password=password
            )
            
            cur = conn.cursor()
            cur.execute(sql)
            rows = cur.fetchall()
            
            columns = [desc[0] for desc in cur.description] if cur.description else []
            
            cur.close()
            conn.close()

            if not columns:
                return {"result": rows}
            
            result = []
            for row in rows:
                result.append(dict(zip(columns, row)))
            return {"columns": columns, "data": result}
        except ImportError:
            raise ValueError("请安装 psycopg2-binary 库")
        except psycopg2.Error as e:
            raise ValueError(f"数据库查询失败：{str(e)}")
        except Exception as e:
            raise ValueError(f"Kingbase 数据获取失败：{str(e)}")

    def _substitute_sql_params(self, sql: str, params: List[dict]) -> str:
        if not params:
            return sql
        
        import re
        from datetime import datetime, date
        
        for param in params:
            param_name = param.get("name", "")
            default_value = param.get("default", "")
            param_type = param.get("type", "string")
            
            if not param_name:
                continue
            
            param_placeholder = f":{param_name}"
            
            if param_placeholder not in sql:
                continue
            
            if default_value:
                if param_type == "date":
                    try:
                        dt = datetime.strptime(default_value, "%Y-%m-%d")
                        sql = sql.replace(param_placeholder, f"'{default_value}'")
                    except:
                        sql = sql.replace(param_placeholder, f"'{default_value}'")
                elif param_type == "integer":
                    sql = sql.replace(param_placeholder, str(default_value))
                else:
                    sql = sql.replace(param_placeholder, f"'{default_value}'")
            else:
                sql = sql.replace(param_placeholder, "NULL")
        
        return sql

    def _substitute_params_in_dict(self, data: dict, params: list) -> dict:
        if not data or not params:
            return data
        
        result = {}
        for key, value in data.items():
            if isinstance(value, str):
                result[key] = self._substitute_sql_params(value, params)
            elif isinstance(value, dict):
                result[key] = self._substitute_params_in_dict(value, params)
            else:
                result[key] = value
        return result

    async def _fetch_postgresql(self, config: dict, runtime_params: Optional[dict] = None) -> dict:
        config = {**config, "port": config.get("port", "5432")}
        return await self._fetch_kingbase(config, runtime_params)

    async def _fetch_mysql(self, config: dict, runtime_params: Optional[dict] = None) -> dict:
        import pymysql
        runtime_params = runtime_params or {}
        try:
            host = config.get("host", "")
            port = config.get("port", "3306")
            database = config.get("database", "")
            user = config.get("user", "")
            password = config.get("password", "")
            sql = config.get("sql", "") or config.get("test_sql", "SELECT 1")
            params = config.get("params", [])
            
            for p in params:
                p_name = p.get("name", "")
                if p_name and p_name in runtime_params:
                    p["default"] = runtime_params[p_name]

            sql = self._substitute_sql_params(sql, params)

            conn = pymysql.connect(
                host=host,
                port=int(port),
                database=database,
                user=user,
                password=password
            )
            
            cur = conn.cursor()
            cur.execute(sql)
            rows = cur.fetchall()
            
            columns = [desc[0] for desc in cur.description] if cur.description else []
            
            cur.close()
            conn.close()

            if not columns:
                return {"result": rows}
            
            result = []
            for row in rows:
                result.append(dict(zip(columns, row)))
            return {"columns": columns, "data": result}
        except ImportError:
            raise ValueError("请安装 pymysql 库")
        except Exception as e:
            raise ValueError(f"MySQL 数据获取失败：{str(e)}")

    async def _fetch_oracle(self, config: dict, runtime_params: Optional[dict] = None) -> dict:
        import cx_Oracle
        runtime_params = runtime_params or {}
        try:
            host = config.get("host", "")
            port = config.get("port", "1521")
            service_name = config.get("service_name", "")
            service = config.get("service", "")
            user = config.get("user", "")
            password = config.get("password", "")
            sql = config.get("sql", "") or config.get("test_sql", "SELECT 1 FROM DUAL")
            params = config.get("params", [])
            
            for p in params:
                p_name = p.get("name", "")
                if p_name and p_name in runtime_params:
                    p["default"] = runtime_params[p_name]

            sql = self._substitute_sql_params(sql, params)

            if service:
                dsn = service
            else:
                dsn = cx_Oracle.makedsn(host, port, service_name=service_name)

            conn = cx_Oracle.connect(user, password, dsn)
            
            cur = conn.cursor()
            cur.execute(sql)
            rows = cur.fetchall()
            
            columns = [desc[0] for desc in cur.description] if cur.description else []
            
            cur.close()
            conn.close()

            if not columns:
                return {"result": rows}
            
            result = []
            for row in rows:
                result.append(dict(zip(columns, row)))
            return {"columns": columns, "data": result}
        except ImportError:
            raise ValueError("请安装 cx-Oracle 库")
        except Exception as e:
            raise ValueError(f"Oracle 数据获取失败：{str(e)}")

    async def _fetch_sql(self, config: dict, runtime_params: Optional[dict] = None, session: AsyncSession = None) -> dict:
        from app.repositories.datasource_repository import DataSourceRepository
        runtime_params = runtime_params or {}
        
        database_id = config.get("database_id")
        if not database_id:
            raise ValueError("SQL 数据源缺少数据库配置")
        
        if not session:
            raise ValueError("SQL 数据源需要数据库会话")
        
        repo = DataSourceRepository(session)
        db_config = await repo.get_by_id(database_id)
        
        if not db_config:
            raise ValueError(f"找不到数据库配置 (ID: {database_id})")
        
        db_type = db_config.type
        db_conn_config = dict(db_config.config or {})
        
        sql = config.get("sql", "") or config.get("test_sql", "SELECT 1")
        params = config.get("params", [])
        
        for p in params:
            p_name = p.get("name", "")
            if p_name and p_name in runtime_params:
                p["default"] = runtime_params[p_name]
        
        sql = self._substitute_sql_params(sql, params)
        
        if db_type == "kingbase":
            return await self._fetch_kingbase({**db_conn_config, "sql": sql})
        elif db_type == "postgresql":
            return await self._fetch_kingbase({**db_conn_config, "sql": sql})
        elif db_type == "mysql":
            return await self._fetch_mysql({**db_conn_config, "sql": sql})
        elif db_type == "oracle":
            return await self._fetch_oracle({**db_conn_config, "sql": sql})
        else:
            raise ValueError(f"不支持的数据库类型: {db_type}")

    def _apply_field_mapping_multi(self, datasource_data: Dict[int, dict], mappings: List[Dict[str, Any]]) -> dict:
        result = {}
        if not datasource_data:
            return result
        
        has_valid_mappings = any(
            mapping.get("source") and mapping.get("target")
            for mapping in mappings
        ) if mappings else False
        
        if not mappings or not has_valid_mappings:
            for ds_id, data in datasource_data.items():
                if isinstance(data, dict):
                    result.update(data)
            return result

        for mapping in mappings:
            ds_id = mapping.get("datasource_id")
            source_field = mapping.get("source", "") or mapping.get("datasource_field", "")
            target_field = mapping.get("target", "") or mapping.get("template_var", "")

            if not target_field or not source_field:
                continue

            data = datasource_data.get(ds_id, {}) if ds_id else {}
            if not data:
                if datasource_data:
                    for _, v in datasource_data.items():
                        data = v
                        break

            if "." in source_field:
                parts = source_field.split(".")
                value = data
                for part in parts:
                    if isinstance(value, dict):
                        value = value.get(part)
                    elif isinstance(value, list) and part.isdigit():
                        idx = int(part)
                        value = value[idx] if idx < len(value) else None
                    else:
                        value = None
                        break
                result[target_field] = value if value is not None else ""
            else:
                value = data.get(source_field, "")
                result[target_field] = value

        return result

    def _generate_charts_for_document(self, chart_configs: List[dict], datasource_data: dict) -> List[dict]:
        if not chart_configs:
            return []

        try:
            for chart_config in chart_configs:
                if chart_config.get("datasource_id"):
                    chart_config = self._enrich_chart_from_datasource(chart_config, datasource_data)
            
            results = chart_service.generate_charts(chart_configs)
            return results
        except Exception as e:
            raise ValueError(f"生成图表失败：{str(e)}")
    
    def _enrich_chart_from_datasource(self, chart_config: dict, datasource_data: dict) -> dict:
        ds_id = chart_config.get("datasource_id")
        if not ds_id or ds_id not in datasource_data:
            return chart_config
        
        data = datasource_data[ds_id]
        x_field = chart_config.get("x_field")
        y_field = chart_config.get("y_field")
        
        if not x_field or not y_field:
            return chart_config
        
        if isinstance(data, list) and len(data) > 0:
            x_data = []
            y_data = []
            for row in data:
                if isinstance(row, dict):
                    x_data.append(str(row.get(x_field, "")))
                    y_data.append(float(row.get(y_field, 0)))
            
            if x_data and y_data:
                chart_config["x_data"] = x_data
                chart_config["y_data"] = y_data
        
        return chart_config

    def _cleanup_charts(self, chart_results: List[dict]):
        for chart_result in chart_results:
            filepath = chart_result.get("filepath")
            if filepath and os.path.exists(filepath):
                try:
                    os.remove(filepath)
                except:
                    pass

    async def _get_datasource_names(self, datasource_ids: List[int], session: AsyncSession) -> str:
        from app.services.datasource_service import data_source_service
        names = []
        for ds_id in datasource_ids:
            ds = await data_source_service.get_datasource(ds_id, session)
            if ds:
                names.append(ds.get("name", f"数据源{ds_id}"))
        return ", ".join(names) if names else "未知"

    async def generate_document(self, mapping_id: int, output_name: Optional[str] = None, runtime_params: Optional[dict] = None, session: AsyncSession = None) -> dict:
        runtime_params = runtime_params or {}
        chart_results = []

        try:
            mapping = await self.mapping_repo.get_by_id(mapping_id, eager_load=True)
            if not mapping:
                raise ValueError("映射配置不存在")

            template_id = mapping.template_id
            datasource_ids = mapping.datasource_ids
            field_mappings = mapping.mappings or []
            chart_configs = mapping.charts

            if not datasource_ids:
                raise ValueError("映射配置中未指定数据源，请至少选择一个数据源")

            template = await self.template_repo.get_by_id(template_id)
            if not template:
                raise ValueError("模板不存在")

            template_file_path = template.file_path
            if not template_file_path:
                raise ValueError("模板路径不存在")
            
            from app.services.template_service import template_service
            local_template_path = await template_service.get_template_file(template_id, session)
            if not local_template_path or not os.path.exists(local_template_path):
                raise ValueError("模板文件不存在或已被删除")
            
            template_file_path = local_template_path

            datasource_data = await self._fetch_multiple_datasources(datasource_ids, session, runtime_params)

            context = self._apply_field_mapping_multi(datasource_data, field_mappings)

            if chart_configs:
                try:
                    chart_results = self._generate_charts_for_document(chart_configs, datasource_data)
                    for chart_result in chart_results:
                        if "error" in chart_result:
                            raise ValueError(f"图表生成失败：{chart_result.get('error')}")
                        chart_id = chart_result.get("id")
                        if chart_id:
                            context[f"chart_{chart_id}"] = chart_result.get("url", "")
                except ValueError:
                    raise
                except Exception as e:
                    raise ValueError(f"生成图表失败：{str(e)}")

            try:
                doc = DocxTemplate(template_file_path)
            except Exception as e:
                raise ValueError(f"加载模板失败：{str(e)}")

            chart_image_map = {}
            for chart_result in chart_results:
                chart_id = chart_result.get("id")
                if chart_id and "filepath" in chart_result:
                    chart_key = f"CHART_PLACEHOLDER_{chart_id}"
                    chart_image_map[chart_key] = chart_result["filepath"]

            try:
                doc.render(context)
            except Exception as e:
                raise ValueError(f"渲染模板失败：{str(e)}")

            temp_rendered = os.path.join(self.output_path, f"_rendered_{uuid.uuid4().hex[:8]}.docx")
            try:
                doc.save(temp_rendered)
            except Exception as e:
                raise ValueError(f"保存渲染中间文件失败：{str(e)}")

            try:
                document = DocxDocument(temp_rendered)
            except Exception as e:
                if os.path.exists(temp_rendered):
                    os.remove(temp_rendered)
                raise ValueError(f"加载渲染文件失败：{str(e)}")

            for paragraph in document.paragraphs:
                for key, img_path in chart_image_map.items():
                    if key in paragraph.text:
                        try:
                            paragraph.clear()
                            run = paragraph.add_run()
                            with open(img_path, 'rb') as f:
                                run.add_picture(f, width=Inches(5))
                        except Exception as e:
                            raise ValueError(f"插入图表到段落失败：{str(e)}")
                        break

            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for key, img_path in chart_image_map.items():
                                if key in paragraph.text:
                                    try:
                                        paragraph.clear()
                                        run = paragraph.add_run()
                                        with open(img_path, 'rb') as f:
                                            run.add_picture(f, width=Inches(4))
                                    except Exception as e:
                                        raise ValueError(f"插入图表到表格失败：{str(e)}")

        except ValueError:
            self._cleanup_charts(chart_results)
            raise
        except Exception as e:
            self._cleanup_charts(chart_results)
            raise ValueError(f"生成文档失败：{str(e)}")

        unique_id = uuid.uuid4().hex[:12]
        original_output_name = output_name
        if output_name:
            safe_output_name = "".join(c for c in output_name if c.isalnum() or c in (' ', '-', '_')).strip()
            if not safe_output_name:
                safe_output_name = "document"
            if not safe_output_name.endswith('.docx'):
                safe_output_name += '.docx'
            filename = f"{safe_output_name}_{unique_id}.docx"
        else:
            filename = f"document_{unique_id}.docx"

        output_file = os.path.join(self.output_path, filename)
        
        from app.services.storage_service import storage_service

        try:
            document.save(output_file)
        except Exception as e:
            raise ValueError(f"保存文档失败：{str(e)}")
        
        with open(output_file, 'rb') as f:
            file_data = f.read()
        
        object_name = await storage_service.upload_file(
            "output",
            filename,
            file_data,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

        if os.path.exists(temp_rendered):
            try:
                os.remove(temp_rendered)
            except Exception:
                pass

        if not os.path.exists(output_file):
            raise ValueError("文档文件生成失败")

        try:
            file_size = os.path.getsize(output_file)
        except Exception:
            file_size = 0

        datasource_names = await self._get_datasource_names(datasource_ids, session)

        history_entry = DocumentHistory(
            mapping_id=mapping_id,
            template_id=template_id,
            datasource_ids=datasource_ids,
            output_filename=original_output_name or filename,
            file_path=object_name,
            status="成功"
        )
        await self.history_repo.add(history_entry)
        await session.commit()

        download_url = await storage_service.get_file_url(object_name)
        
        return {
            "id": history_entry.id,
            "name": history_entry.output_filename,
            "template": template.name,
            "datasource": datasource_names,
            "status": history_entry.status,
            "file_path": object_name,
            "download_url": download_url,
            "file_size": self._format_file_size(file_size),
            "create_time": history_entry.create_time.strftime("%Y-%m-%d %H:%M:%S") if history_entry.create_time else ""
        }

    def _format_file_size(self, size: int) -> str:
        for unit in ['B', 'KB', 'MB', 'GB']:
            if size < 1024:
                return f"{size:.1f}{unit}"
            size /= 1024
        return f"{size:.1f}TB"

    async def get_history(self, page: int = 1, page_size: int = 10, search: Optional[str] = None, status: Optional[str] = None, session: AsyncSession = None) -> dict:
        history_items, total = await self.history_repo.get_paginated(
            page, page_size, search, status
        )

        items = []
        for h in history_items:
            file_size = await storage_service.get_file_size(h.file_path) if h.file_path else 0
            items.append({
                "id": h.id,
                "name": h.output_filename,
                "template": "",
                "datasource": "",
                "status": h.status,
                "file_size": self._format_file_size(file_size),
                "create_time": h.create_time.strftime("%Y-%m-%d %H:%M:%S") if h.create_time else ""
            })

        return {
            "total": total,
            "page": page,
            "page_size": page_size,
            "data": items
        }

    async def get_history_item(self, history_id: int, session: AsyncSession) -> Optional[dict]:
        h = await self.history_repo.get_by_id(history_id)
        if not h:
            return None
        return {
            "id": h.id,
            "name": h.output_filename,
            "template": "",
            "template_id": h.template_id,
            "datasource": "",
            "datasource_ids": h.datasource_ids or [],
            "mapping_id": h.mapping_id,
            "status": h.status,
            "file_size": self._format_file_size(
                await storage_service.get_file_size(h.file_path) if h.file_path else 0
            ),
            "create_time": h.create_time.strftime("%Y-%m-%d %H:%M:%S") if h.create_time else ""
        }

    async def download_document(self, history_id: int, session: AsyncSession) -> Optional[str]:
        from app.services.storage_service import storage_service
        h = await self.history_repo.get_by_id(history_id)
        if h and h.file_path:
            try:
                data = await storage_service.download_file(h.file_path)
                temp_path = f"/tmp/{h.output_filename}"
                with open(temp_path, 'wb') as f:
                    f.write(data)
                return temp_path
            except:
                pass
        return None

    async def delete_history_item(self, history_id: int, session: AsyncSession) -> bool:
        from app.services.storage_service import storage_service
        h = await self.history_repo.get_by_id(history_id)
        if not h:
            return False

        if h.file_path:
            try:
                await storage_service.delete_file(h.file_path)
            except:
                pass

        await self.history_repo.delete(history_id)
        await session.commit()
        return True

    async def regenerate_document(self, history_id: int, session: AsyncSession) -> dict:
        history_item = await self.get_history_item(history_id, session)
        if not history_item:
            raise ValueError("历史记录不存在")

        mapping_id = history_item.get("mapping_id")
        if not mapping_id:
            raise ValueError("无法重新生成：缺少映射配置 ID")

        return await self.generate_document(mapping_id, history_item.get("name"), session)


document_generation_service = None
