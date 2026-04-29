import os
import re
import json
import uuid
from datetime import datetime
from docx import Document
from typing import List, Optional
from io import BytesIO
from sqlalchemy.ext.asyncio import AsyncSession
from app.core.config import settings
from app.models import Template
from app.repositories.template_repository import TemplateRepository
from app.services.storage_service import storage_service


class TemplateService:
    def __init__(self):
        self.storage_prefix = "templates"

    def extract_variables_from_bytes(self, file_data: bytes) -> List[str]:
        variables = set()
        try:
            with BytesIO(file_data) as bio:
                doc = Document(bio)
                for paragraph in doc.paragraphs:
                    vars_in_text = self._extract_vars_from_text(paragraph.text)
                    variables.update(vars_in_text)

                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                vars_in_text = self._extract_vars_from_text(paragraph.text)
                                variables.update(vars_in_text)
        except Exception as e:
            print(f"Error extracting variables: {e}")

        return sorted(list(variables))

    def _extract_vars_from_text(self, text: str) -> List[str]:
        pattern = r'\{\{([^}]+)\}\}'
        matches = re.findall(pattern, text)
        return [m.strip() for m in matches]

    async def upload_template(self, file, name: str, category: str, session: AsyncSession) -> dict:
        content_type = file.content_type or ''

        if content_type not in ['application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'application/octet-stream', '']:
            raise ValueError("只支持 .docx 格式的文件")

        file_size = file.size or 0
        if file_size > settings.MAX_TEMPLATE_SIZE:
            raise ValueError(f"文件大小不能超过 {settings.MAX_TEMPLATE_SIZE / 1024 / 1024}MB")

        original_filename = file.filename or 'template.docx'
        
        content = await file.read()
        
        object_name = await storage_service.upload_file(
            self.storage_prefix,
            original_filename,
            content,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        
        variables = self.extract_variables_from_bytes(content)
        
        # Use repository to get max ID
        repo = TemplateRepository(session)
        all_templates = await repo.get_all()
        max_id = max((t.id for t in all_templates), default=0)
        new_id = max_id + 1

        template = Template(
            id=new_id,
            name=name,
            category=category,
            file_name=original_filename,
            file_path=object_name,
            file_size=file_size,
            variables=variables
        )
        session.add(template)
        await session.commit()
        await session.refresh(template)

        return {
            "id": template.id,
            "name": name,
            "category": category,
            "file_name": original_filename,
            "file_size": file_size,
            "variables": variables,
            "create_time": template.create_time.strftime("%Y-%m-%d %H:%M:%S") if template.create_time else ""
        }

    async def get_templates(self, session: AsyncSession) -> List[dict]:
        repo = TemplateRepository(session)
        templates = await repo.get_all()
        return [
            {
                "id": t.id,
                "name": t.name,
                "category": t.category,
                "file_name": t.file_name,
                "file_size": t.file_size,
                "variables": t.variables or [],
                "create_time": t.create_time.strftime("%Y-%m-%d %H:%M:%S") if t.create_time else ""
            }
            for t in templates
        ]

    async def get_templates_paginated(self, page: int = 1, page_size: int = 10, session: AsyncSession = None) -> dict:
        if session is None:
            from app.models import async_session
            async with async_session() as session:
                return await self._get_templates_paginated(page, page_size, session)
        return await self._get_templates_paginated(page, page_size, session)

    async def _get_templates_paginated(self, page: int, page_size: int, session: AsyncSession) -> dict:
        repo = TemplateRepository(session)
        templates, total = await repo.get_paginated(page, page_size)

        items = [
            {
                "id": t.id,
                "name": t.name,
                "category": t.category,
                "file_name": t.file_name,
                "file_size": t.file_size,
                "variables": t.variables or [],
                "create_time": t.create_time.strftime("%Y-%m-%d %H:%M:%S") if t.create_time else ""
            }
            for t in templates
        ]

        return {
            "total": total,
            "page": page,
            "page_size": page_size,
            "data": items
        }

    async def get_template(self, template_id: int, session: AsyncSession) -> Optional[dict]:
        repo = TemplateRepository(session)
        t = await repo.get_by_id(template_id)
        if not t:
            return None
        return {
            "id": t.id,
            "name": t.name,
            "category": t.category,
            "file_name": t.file_name,
            "file_size": t.file_size,
            "variables": t.variables or [],
            "create_time": t.create_time.strftime("%Y-%m-%d %H:%M:%S") if t.create_time else ""
        }

    async def delete_template(self, template_id: int, session: AsyncSession) -> bool:
        repo = TemplateRepository(session)
        t = await repo.get_by_id(template_id)
        if not t:
            return False

        if t.file_path:
            try:
                await storage_service.delete_file(t.file_path)
            except:
                pass

        await session.delete(t)
        await session.commit()
        return True

    async def get_template_file(self, template_id: int, session: AsyncSession) -> Optional[str]:
        repo = TemplateRepository(session)
        t = await repo.get_by_id(template_id)
        if t and t.file_path:
            try:
                data = await storage_service.download_file(t.file_path)
                temp_path = f"/tmp/{t.file_name}"
                with open(temp_path, 'wb') as f:
                    f.write(data)
                return temp_path
            except:
                return None
        return None
    
    async def get_template_download_url(self, template_id: int, session: AsyncSession) -> Optional[str]:
        repo = TemplateRepository(session)
        t = await repo.get_by_id(template_id)
        if t and t.file_path:
            try:
                return await storage_service.get_file_url(t.file_path)
            except:
                return None
        return None


template_service = TemplateService()